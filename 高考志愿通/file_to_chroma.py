import openpyxl
import os
import sqlite3  # 新增
from langchain.vectorstores import Chroma
from langchain.embeddings import OpenAIEmbeddings
from langchain.schema import Document
from dotenv import load_dotenv

# 加载环境变量
load_dotenv()
os.environ["OPENAI_API_KEY"] = os.getenv('OPENAI_API_KEY')

# 使用 OpenAI 的嵌入模型
embeddings = OpenAIEmbeddings()

def xlsx_parse(file):
    try:
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.active

        # 获取列标题
        headers = [cell.value for cell in sheet[1]]
        rows = []

        # 读取每一行的数据
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if any(cell is not None for cell in row):  # 确保至少有一个单元格有内容
                row_data = {headers[i]: str(cell) if cell is not None else "" for i, cell in enumerate(row)}
                rows.append(row_data)

        return rows
    except Exception as e:
        print(f"解析 {file} 时出错: {e}")
        return []

def process_file(file_path, vectorstore, cursor):  # 新增参数 cursor
    rows = xlsx_parse(file_path)
    if not rows:
        return

    # 获取列名
    headers = rows[0].keys()

    # 生成合法的表名（移除非法字符）
    table_name = os.path.splitext(os.path.basename(file_path))[0]
    table_name = table_name.replace(' ', '_').replace('-', '_')

    # 创建表（如果不存在）
    create_table_sql = f"CREATE TABLE IF NOT EXISTS [{table_name}] ({', '.join(f'[{col}] TEXT' for col in headers)});"
    cursor.execute(create_table_sql)

    # 插入数据
    for row in rows:
        placeholders = ', '.join('?' for _ in headers)
        insert_sql = f"INSERT INTO [{table_name}] ({', '.join(f'[{col}]' for col in headers)}) VALUES ({placeholders});"
        values = [row[col] for col in headers]
        cursor.execute(insert_sql, values)

    # 继续原有的向量库处理
    documents = []

    for i, row in enumerate(rows):
        row_text = " ".join(f"{key}: {value}" for key, value in row.items())
        # 将每一行转换为 Document 对象
        document = Document(page_content=row_text, metadata=row)
        documents.append(document)
        print(f"添加文档内容: {row_text}")  # 打印文档内容以进行调试

    # 添加到 LangChain 的 Chroma 向量库
    try:
        vectorstore.add_documents(documents=documents)
        print(f"{os.path.basename(file_path)} 数据添加成功")
    except Exception as e:
        print(f"添加到 Chroma 向量库时出错: {e}")

def run(all_file_path):
    # 创建 ChromaDB 的持久化路径
    persist_directory = "/Users/leo/研究生/实习/北大青鸟/项目制作/高考政策通/chromadb_rag"

    # 使用 LangChain 和 OpenAIEmbeddings 创建一个向量数据库
    vectorstore = Chroma(
        embedding_function=embeddings,
        persist_directory=persist_directory,
    )

    # 连接到 SQLite 数据库（或创建数据库）
    db_path = "/Users/leo/研究生/实习/北大青鸟/项目制作/高考政策通/database.db"
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    # 处理所有 Excel 文件
    file_names = os.listdir(all_file_path)
    for file_name in file_names:
        file_path = os.path.join(all_file_path, file_name)
        if file_path.endswith(".xlsx"):
            print(f"正在处理文件: {file_path}")
            process_file(file_path, vectorstore, cursor)

    # 提交并关闭数据库连接
    conn.commit()
    conn.close()

    # 持久化存储
    vectorstore.persist()

    return "向量化完成！"

if __name__ == '__main__':
    all_file_path = "/Users/leo/研究生/实习/北大青鸟/项目制作/高考政策通/2016-2021高考数据/"
    result = run(all_file_path)
    print(result)
